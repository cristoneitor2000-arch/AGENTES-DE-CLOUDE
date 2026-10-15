"""extraccion.py – Excel de data extraction (Arial 10, negro, centrado).

Hojas
  DATOS       una fila por objeto creado (con su handle en el DXF de salida)
  RESUMEN     por elemento y capa, con COUNTIFS/SUMIFS sobre DATOS (fórmulas clásicas)
  REVISION    marcas de la capa REVISION con su motivo
  PARAMETROS  entradas del proceso (archivo, unidades, espesores, tolerancias, sha256)
Los valores van sin redondear; el número de decimales es solo formato de celda.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

F = Font(name="Arial", size=10, color="000000")
FB = Font(name="Arial", size=10, color="000000", bold=True)
C = Alignment(horizontal="center", vertical="center", wrap_text=True)
S = Side(style="thin", color="000000")
B = Border(left=S, right=S, top=S, bottom=S)
GR = PatternFill("solid", fgColor="D9D9D9")

COLS_DATOS = [("ID", 10), ("Elemento", 18), ("Capa", 44), ("Espesor de muro (m)", 12), ("Tipo de cimiento", 14),
              ("Ancho de cimiento (m)", 12), ("Peralte de cimiento (m)", 12), ("Longitud", 12), ("Unidad longitud", 9),
              ("Altura", 10), ("Unidad altura", 9), ("Área", 12), ("Unidad área", 9), ("Handle", 10),
              ("X inicio (m)", 12), ("Y inicio (m)", 12), ("X fin (m)", 12), ("Y fin (m)", 12), ("Observaciones", 50)]


def _hoja(wb, nombre, cols):
    ws = wb.create_sheet(nombre)
    for j, (t, w) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=j, value=t)
        c.font, c.alignment, c.border, c.fill = FB, C, B, GR
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    return ws


def _fila(ws, r, valores, formatos=None):
    for j, v in enumerate(valores, start=1):
        c = ws.cell(row=r, column=j, value=v)
        c.font, c.alignment, c.border = F, C, B
        if formatos and j in formatos and v not in (None, ""):
            c.number_format = formatos[j]


def escribir(ruta, elementos, revisiones, parametros):
    wb = Workbook()
    wb.remove(wb.active)
    ws = _hoja(wb, "DATOS", COLS_DATOS)
    fmt = {4: "0.00", 6: "0.00", 7: "0.00", 8: "0.000", 10: "0.00", 12: "0.000", 15: "0.000", 16: "0.000", 17: "0.000", 18: "0.000"}
    for i, e in enumerate(elementos, start=2):
        alt = e.get("altura")
        area = f'=IF(J{i}="","",H{i}*J{i})' if e["elemento"] == "Muro" else None
        _fila(ws, i, [e["id"], e["elemento"], e["capa"], e.get("esp"), e.get("tipo_cim"), e.get("ancho"), e.get("peralte"),
                      e["largo"], "m", alt, "m" if alt is not None else None, area, "m²" if e["elemento"] == "Muro" else None,
                      e.get("handle"), e["p0"][0], e["p0"][1], e["p1"][0], e["p1"][1], e.get("obs", "")], fmt)
    n = len(elementos) + 1

    # RESUMEN
    wr = _hoja(wb, "RESUMEN", [("Elemento", 18), ("Capa", 44), ("Cantidad de objetos", 12), ("Unidad cantidad", 9),
                               ("Longitud", 12), ("Unidad longitud", 9), ("Área", 12), ("Unidad área", 9)])
    claves = []
    for e in elementos:
        k = (e["elemento"], e["capa"])
        if k not in claves:
            claves.append(k)
    orden = {"Cimiento corrido": 0, "Muro": 1, "Vano puerta": 2}
    claves.sort(key=lambda k: (orden.get(k[0], 9), k[1]))
    r = 2
    rng = lambda col: f"DATOS!${col}$2:${col}${max(n, 2)}"
    for el, cp in claves:
        _fila(wr, r, [el, cp, f'=COUNTIFS({rng("B")},A{r},{rng("C")},B{r})', "u",
                      f'=SUMIFS({rng("H")},{rng("B")},A{r},{rng("C")},B{r})', "m",
                      f'=SUMIFS({rng("L")},{rng("B")},A{r},{rng("C")},B{r})' if el == "Muro" else None,
                      "m²" if el == "Muro" else None], {5: "0.000", 7: "0.000"})
        r += 1
    r += 1
    for el in sorted({k[0] for k in claves}, key=lambda x: orden.get(x, 9)):
        _fila(wr, r, [f"Total {el}", "(todas las capas)", f'=COUNTIFS({rng("B")},"{el}")', "u",
                      f'=SUMIFS({rng("H")},{rng("B")},"{el}")', "m",
                      f'=SUMIFS({rng("L")},{rng("B")},"{el}")' if el == "Muro" else None, "m²" if el == "Muro" else None],
              {5: "0.000", 7: "0.000"})
        for j in range(1, 9):
            wr.cell(row=r, column=j).font = FB
        r += 1

    # REVISION
    wv = _hoja(wb, "REVISION", [("Marca", 8), ("Motivo", 24), ("X (m)", 12), ("Y (m)", 12), ("Handles de origen", 30), ("Detalle", 70)])
    for i, rv in enumerate(revisiones, start=1):
        _fila(wv, i + 1, [f"R{i:03d}", rv.motivo, rv.x, rv.y, ", ".join(rv.handles), rv.detalle], {3: "0.000", 4: "0.000"})
    nv = len(revisiones) + 1
    col0 = 8
    wv.cell(row=1, column=col0, value="Motivo").font = FB
    wv.cell(row=1, column=col0 + 1, value="Marcas").font = FB
    for j in (col0, col0 + 1):
        c = wv.cell(row=1, column=j)
        c.alignment, c.border, c.fill = C, B, GR
    wv.column_dimensions[get_column_letter(col0)].width = 24
    for i, m in enumerate(sorted({rv.motivo for rv in revisiones}), start=2):
        _fila_en(wv, i, col0, [m, f'=COUNTIF($B$2:$B${max(nv, 2)},{get_column_letter(col0)}{i})'])

    # PARAMETROS
    wp = _hoja(wb, "PARAMETROS", [("Parámetro", 34), ("Valor", 70), ("Unidad", 10)])
    for i, (k, v, u) in enumerate(parametros, start=2):
        _fila(wp, i, [k, v, u])
    wb.save(ruta)
    return ruta


def _fila_en(ws, r, c0, valores):
    for j, v in enumerate(valores):
        c = ws.cell(row=r, column=c0 + j, value=v)
        c.font, c.alignment, c.border = F, C, B
