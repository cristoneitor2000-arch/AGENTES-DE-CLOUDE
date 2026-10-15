"""Proceso completo: DXF de salida, QA de la skill, Excel, reporte y original intacto."""
import copy
import os
import subprocess
import sys

import ezdxf
import openpyxl
import pytest

import escritor
import lector
import pipeline

QA = os.path.expanduser("~/.claude/skills/synced")


@pytest.fixture(scope="module")
def info(plano, conf, tipos_prueba, tmp_path_factory):
    conv, cfg = conf
    return pipeline.cuantificar(plano, ["A-MURO"], [.10, .15, .20], "N1", 2.80, None, conv, cfg,
                                log=lambda m: None, carpeta_salida=str(tmp_path_factory.mktemp("salida")),
                                tipos_cimiento=tipos_prueba)


def test_original_intacto(info):
    assert info["sha_ini"] == info["sha_fin"]


def test_archivos(info):
    s = info["salida"]
    assert os.path.basename(s).startswith("SALIDA_prueba_muros_")
    for f in ("prueba_muros - CUANT.dxf", "prueba_muros - COPIA ORIGINAL.dxf", "DataExtraction_prueba_muros.xlsx", "reporte.md"):
        assert os.path.isfile(os.path.join(s, f)), f


def test_dxf_salida(info):
    doc = ezdxf.readfile(info["dxf"])
    assert doc.audit().has_errors is False
    msp = doc.modelspace()
    L = doc.layers
    assert L.get("00").is_locked() and L.get("00").color == 8
    assert L.get("A-MURO").is_locked()
    # todo lo del plano original quedó en 00
    orig = ezdxf.readfile(info["original"])
    assert sum(1 for e in msp if e.dxf.layer == "00") == len(orig.modelspace())
    # capas de cuantificación
    mur = L.get("N1 Muro de 15cm")
    assert mur.color == 1 and abs(mur.transparency - 0.5) < 0.01 and not mur.is_locked()
    assert L.get("N1 Muro de 15cm - Vano puerta").color == 42
    # en capas de muro solo INSERT del bloque estático, escala X = largo
    largos = {}
    for e in msp:
        if " Muro de " in e.dxf.layer:
            assert e.dxftype() == "INSERT" and e.dxf.name.startswith("xMuro ")
            at = {a.dxf.tag: a.dxf.text for a in e.attribs}
            assert abs(float(at["LARGO"]) - e.dxf.xscale) < 1e-4
            if "Vano" not in e.dxf.layer:
                assert abs(float(at["AREA"]) - e.dxf.xscale * 2.80) < 1e-3
            largos[e.dxf.layer] = largos.get(e.dxf.layer, 0) + e.dxf.xscale
    assert largos["N1 Muro de 15cm"] == pytest.approx(41.725, abs=1e-9)
    # cimiento: polilínea abierta con ancho global
    cim = [e for e in msp if e.dxf.layer.startswith("Cimiento corrido")]
    assert cim and all(e.dxftype() == "LWPOLYLINE" and not e.closed for e in cim)
    assert all(abs(e.dxf.const_width - (0.40 if "0.40" in e.dxf.layer else 0.50)) < 1e-12 for e in cim)
    # REVISION: círculo rojo + texto por marca
    rv = [e for e in msp if e.dxf.layer == "REVISION"]
    assert sum(1 for e in rv if e.dxftype() == "CIRCLE") == 4 == sum(1 for e in rv if e.dxftype() == "TEXT")   # 3 + esquina L
    assert all(e.dxf.color == 1 for e in rv)
    # grupos
    g = escritor.leer_grupos(doc)
    assert set(g["CUANTIFICACION > MUROS > MUROS N1"]) == {"N1 Muro de 10cm", "N1 Muro de 15cm", "N1 Muro de 20cm", "N1 Muro de 15cm - Vano puerta"}
    assert set(g["CUANTIFICACION > CIMIENTO CORRIDO"]) == {"Cimiento corrido CC1 de 0.40 x 0.20 m", "Cimiento corrido MC1 de 0.50 x 0.20 m"}
    # cimiento al fondo del orden de dibujo
    orden = dict(msp.get_redraw_order())
    assert all(e.dxf.handle in orden for e in cim)


def test_qa_dxf_de_la_skill(info, tmp_path):
    rutas = [os.path.join(dp, f) for dp, _, fs in os.walk(QA) for f in fs if f == "qa_dxf.py"]
    if not rutas:
        pytest.skip("qa_dxf.py no disponible en este equipo")
    p = subprocess.run([sys.executable, rutas[0], info["dxf"], str(tmp_path / "qa.xlsx")], capture_output=True, text=True)
    assert p.returncode == 0, p.stdout + p.stderr   # 0 = sin ERROR
    wb = openpyxl.load_workbook(tmp_path / "qa.xlsx")
    reglas = {r[1] for r in wb["HALLAZGOS"].iter_rows(min_row=2, values_only=True)}
    assert not any(r.startswith(("R02", "R03", "R06", "R07", "R10")) for r in reglas), reglas


def test_excel(info):
    wb = openpyxl.load_workbook(info["xlsx"])
    assert wb.sheetnames == ["DATOS", "RESUMEN", "REVISION", "PARAMETROS"]
    ws = wb["DATOS"]
    hdr = [c.value for c in ws[1]]
    for col in ("Elemento", "Capa", "Espesor de muro (m)", "Tipo de cimiento", "Longitud", "Unidad longitud", "Handle", "Observaciones"):
        assert col in hdr
    filas = list(ws.iter_rows(min_row=2, values_only=True))
    ih, il = hdr.index("Handle"), hdr.index("Longitud")
    doc = ezdxf.readfile(info["dxf"])
    for f in filas:
        assert doc.entitydb.get(f[ih]) is not None, f   # cada fila apunta a un objeto real del DXF
    assert sum(f[il] for f in filas if f[1] == "Muro") == pytest.approx(41.725 + 6 + 2.95, abs=1e-9)
    c = ws["A2"]
    assert c.font.name == "Arial" and c.font.sz == 10 and c.alignment.horizontal == "center"


def test_reporte(info):
    txt = open(os.path.join(info["salida"], "reporte.md"), encoding="utf-8").read()
    for s in ("Segmentos analizados", "% reconocido por longitud", "Marcas en REVISION por motivo", "Tiempo de proceso", "sin cambios"):
        assert s in txt


def test_sin_tabla_de_tipos_todo_el_cimiento_a_revision(plano, conf, tmp_path):
    conv, cfg = conf
    assert conv["cimiento"]["tipos"] == {}   # producción: la tabla se carga por proyecto en la interfaz
    info = pipeline.cuantificar(plano, ["A-MURO"], [.10, .15, .20], "N1", None, None, conv, cfg, log=lambda m: None,
                                carpeta_salida=str(tmp_path))
    assert not any(e["elemento"] == "Cimiento corrido" for e in info["elementos"])
    assert sum(1 for r in info["revisiones"] if r.motivo == "cimiento sin etiqueta") == 11   # 12 ejes menos el tabique de 0.10
    doc = ezdxf.readfile(info["dxf"])
    at = [{a.dxf.tag: a.dxf.text for a in e.attribs} for e in doc.modelspace().query("INSERT") if "Muro de" in e.dxf.layer]
    assert all(a["ALTO"] == "" and a["AREA"] == "" and a["LARGO"] for a in at)   # sin altura: solo LARGO


def test_dwg_sin_oda(info):
    assert any("ODA" in a for a in info["avisos"])
    assert not any(f.endswith(".dwg") for f in info["archivos"])


def test_emplantillado_completo(plano_bloques, conf, tipos_prueba, tmp_path):
    conv, cfg = conf
    info = pipeline.cuantificar(plano_bloques, ["BLK", "TABIQUE"], [.10, .15, .20], "SOT", None, None, conv, cfg,
                                log=lambda m: None, carpeta_salida=str(tmp_path), capas_columnas=["COL"], capas_zapatas=["ZAP"],
                                tipos_cimiento=tipos_prueba, espesor_por_capa={"TABIQUE": 0.10})
    assert info["revisiones"] == []
    doc = ezdxf.readfile(info["dxf"])
    L = doc.layers
    assert L.get("SOT Muro de contencion de 30cm").color == 6
    assert L.get("Cimiento corrido CC1 de 0.40 x 0.20 m (dentro de zapata)").color == 52
    g = escritor.leer_grupos(doc)
    assert g["CUANTIFICACION > MURO DE CONTENCION"] == ["SOT Muro de contencion de 30cm"]
    assert "Cimiento corrido CC1 de 0.40 x 0.20 m (dentro de zapata)" in g["CUANTIFICACION > CIMIENTO CORRIDO"]
    tot = {}
    for e in info["elementos"]:
        tot[e["capa"]] = tot.get(e["capa"], 0) + e["largo"]
    assert tot["SOT Muro de contencion de 30cm"] == pytest.approx(6.2)
    assert tot["SOT Muro de 10cm"] == pytest.approx(2.0)
    assert tot["SOT Muro de 15cm"] == pytest.approx(2.0 + 4.3 + 4.3 + 3.2)
    assert tot["SOT Muro de 15cm - Vano puerta"] == pytest.approx(0.9)
    assert tot["Cimiento corrido CC1 de 0.40 x 0.20 m (dentro de zapata)"] == pytest.approx(2.0)
    rutas = [os.path.join(dp, f) for dp, _, fs in os.walk(QA) for f in fs if f == "qa_dxf.py"]
    if rutas:
        p = subprocess.run([sys.executable, rutas[0], info["dxf"], str(tmp_path / "qa.xlsx")], capture_output=True, text=True)
        assert p.returncode == 0, p.stdout + p.stderr
