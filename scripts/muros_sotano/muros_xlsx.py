import sys, json, math
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
D=json.load(open(sys.argv[1])); out=sys.argv[2]
ALL,V,H,CHK=D["ALL"],D["V"],D["H"],D["chk"]
F=Font(name="Arial",size=10); FB=Font(name="Arial",size=10,bold=True); FT=Font(name="Arial",size=12,bold=True); FI=Font(name="Arial",size=9,italic=True,color="595959")
C=Alignment(horizontal="center",vertical="center",wrap_text=True); LF=Alignment(horizontal="left",vertical="center",wrap_text=True)
HDR=PatternFill("solid",fgColor="D9D9D9"); GRP=PatternFill("solid",fgColor="F2F2F2"); INP=PatternFill("solid",fgColor="FFF2CC")
th=Side(style="thin",color="BFBFBF"); B=Border(top=th,bottom=th,left=th,right=th)
SUB="Proyecto: Sótano – Plano S-1 (DXF) | Fecha: 2026-10-12 | Unidades: m, m²"
wb=Workbook()
def sheet(ws,t,h,wd,sub=SUB):
    ws.sheet_view.showGridLines=False; ws["A1"]=t; ws["A1"].font=FT; ws["A2"]=sub; ws["A2"].font=FI
    for i,(x,w) in enumerate(zip(h,wd),1):
        c=ws.cell(row=4,column=i,value=x); c.font=FB; c.fill=HDR; c.alignment=C; c.border=B; ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes="A5"; return 5
def row(ws,r,vals,bold=False,fill=None,fmt=None,left=()):
    for i,v in enumerate(vals,1):
        c=ws.cell(row=r,column=i,value=v); c.font=FB if bold else F; c.alignment=LF if i in left else C; c.border=B
        if fill: c.fill=fill
        if fmt and i in fmt: c.number_format=fmt[i]
TIPOS=["Muro block 0.20 que carga","Muro block 0.15 que carga","Muro block 0.15 que no carga","Muro MC1 block doble hilada 0.15+0.15"]
# Parámetros
wp=wb.active; wp.title="Parámetros"
r=sheet(wp,"PARÁMETROS",["Parámetro","Valor","Unidad","Fuente / estado"],[44,12,10,60],SUB+" | Celdas amarillas = pendiente")
P={}
for n,u,f in [("Altura total muros de edificación (sótano)","m","PENDIENTE – no hay sección/elevación en el DXF"),("Altura total muro MC1","m","PENDIENTE – detalle MC1 (xref)"),("Altura de vanos de puerta","m","PENDIENTE – plano arquitectónico")]:
    row(wp,r,[n,None,u,f],fmt={2:"0.00"},left=(1,4)); wp.cell(row=r,column=2).fill=INP; P[n]=f"Parámetros!$B${r}"; r+=1
HE=P["Altura total muros de edificación (sótano)"]; HM=P["Altura total muro MC1"]; HV=P["Altura de vanos de puerta"]
# Muros (detalle)
wm=wb.create_sheet("Muros")
r=sheet(wm,"MUROS – EXTRACCIÓN POR TRAMO",["ID","Grupo","Tipo","Etiqueta plano","Espesor","Longitud eje","Altura total","Área bruta","Vano","Inicio X,Y","Fin X,Y"],[13,16,34,10,9,11,10,11,8,15,15])
wm.cell(row=3,column=5,value="m").font=FI; wm.cell(row=3,column=6,value="m").font=FI; wm.cell(row=3,column=7,value="m").font=FI; wm.cell(row=3,column=8,value="m²").font=FI
m0=r
for t in TIPOS:
    for w in sorted([w for w in ALL if w["tipo"]==t],key=lambda w:w["wid"]):
        h=HM if w["grupo"]=="MC" else HE
        row(wm,r,[w["wid"],"Muro de contención" if w["grupo"]=="MC" else "Edificación",w["tipo"],w.get("etq",""),round(w["e"],2),round(w["L"],3),f'=IF({h}="","",{h})',f'=IF(G{r}="","",F{r}*G{r})',w.get("vano",""),"%.2f, %.2f"%tuple(w["p"]),"%.2f, %.2f"%tuple(w["q"])],fmt={5:"0.00",6:"0.00",7:"0.00",8:"0.00"},left=(3,))
        r+=1
m1=r-1
row(wm,r,["TOTAL","","","","",f"=SUM(F{m0}:F{m1})","",f'=IF(COUNT(H{m0}:H{m1})<COUNTA(A{m0}:A{m1}),"pendiente",SUM(H{m0}:H{m1}))',"","",""],bold=True,fill=GRP,fmt={6:"0.00",8:"0.00"})
# Vanos
wv=wb.create_sheet("Vanos")
r=sheet(wv,"VANOS – DEDUCCIÓN",["ID","Muro","Tipo de muro","Ancho","Alto","Área a deducir","Centro X,Y","Observaciones"],[8,10,34,9,9,13,15,46])
v0=r
for v in V:
    a=next(w for w in ALL if w["wid"]==v["a"])
    row(wv,r,[v["id"],v["a"],a["tipo"],v["ancho"],f'=IF({HV}="","",{HV})',f'=IF(E{r}="","",D{r}*E{r})',"%.2f, %.2f"%tuple(v["c"]),"Hueco colineal entre %s y %s; muro trazado a altura total sobre el vano"%(v["a"],v["b"])],fmt={4:"0.00",5:"0.00",6:"0.00"},left=(3,8)); r+=1
v1=r-1
row(wv,r,["TOTAL","","",f"=SUM(D{v0}:D{v1})","",f'=IF(COUNT(F{v0}:F{v1})<COUNTA(A{v0}:A{v1}),"pendiente",SUM(F{v0}:F{v1}))',"",""],bold=True,fill=GRP,fmt={4:"0.00",6:"0.00"})
# Resumen
wr=wb.create_sheet("Resumen",0)
r=sheet(wr,"RESUMEN DE CUANTIFICACIÓN – MUROS SÓTANO",["No.","Elemento","Espesor","Longitud","Altura total","Área bruta","(−) Vanos","Área neta","Observaciones"],[6,36,9,11,10,11,10,11,46])
for i,u in [(4,"ml"),(5,"m"),(6,"m²"),(7,"m²"),(8,"m²")]: wr.cell(row=3,column=i,value=u).font=FI; wr.cell(row=3,column=i).alignment=C
def grp(t):
    global r
    for i in range(1,10): c=wr.cell(row=r,column=i); c.fill=GRP; c.border=B
    wr.cell(row=r,column=1,value=t).font=FB; wr.merge_cells(start_row=r,start_column=1,end_row=r,end_column=9); wr.cell(row=r,column=1).alignment=LF; r+=1
n=0
grp("A. MUROS DE EDIFICACIÓN")
for t in TIPOS:
    if t.startswith("Muro MC1"):
        grp("B. MURO DE CONTENCIÓN MC1 (grupo independiente)")
    n+=1; h=HM if t.startswith("Muro MC1") else HE
    e={"Muro block 0.20 que carga":0.20,"Muro block 0.15 que carga":0.15,"Muro block 0.15 que no carga":0.15}.get(t,0.30)
    obs={"Muro block 0.15 que no carga":"Incluye M-40, M-41, M-42 y M-49 (11.63 ml) paralelos al MC1 → RFI","Muro MC1 block doble hilada 0.15+0.15":"Excluye huecos de columnas C; ver Pendientes"}.get(t,"")
    row(wr,r,[n,t,e,f'=SUMIF(Muros!$C${m0}:$C${m1},B{r},Muros!$F${m0}:$F${m1})',f'=IF({h}="","pendiente",{h})',f'=IF({h}="","pendiente",D{r}*{h})',
              f'=IF(COUNTIF(Vanos!$C${v0}:$C${v1},B{r})=0,0,IF({HV}="","pendiente",SUMIF(Vanos!$C${v0}:$C${v1},B{r},Vanos!$D${v0}:$D${v1})*{HV}))',
              f'=IF(OR(F{r}="pendiente",G{r}="pendiente"),"pendiente",F{r}-G{r})',obs],fmt={3:"0.00",4:"0.00",5:"0.00",6:"0.00",7:"0.00",8:"0.00"},left=(2,9)); r+=1
r+=1
for txt in ["Longitud = eje de cada tramo de block dibujado en planta; en vanos el muro corre continuo a ALTURA TOTAL y el área del vano se deduce aparte (hoja Vanos).",
            "Columnas, mochetas y huecos de 0.10–0.85 m entre tramos no se cuentan como muro (16 huecos, ver Pendientes)."]:
    wr.cell(row=r,column=1,value=txt).font=FI; wr.merge_cells(start_row=r,start_column=1,end_row=r,end_column=9); r+=1
# Pendientes
wq=wb.create_sheet("Pendientes")
r=sheet(wq,"PENDIENTES, RFI Y SUPUESTOS",["No.","Tipo","Descripción","Impacto"],[6,12,92,36])
hs=", ".join(f"{h['a']}→{h['b']} {h['g']:.2f}" for h in H)
items=[("Dato faltante","Alturas: el DXF es una planta de cimentación sin secciones ni elevaciones y el xref Xref_Muros_Sotano no está incrustado. Las áreas quedan con fórmula y se calculan al ingresar alturas en Parámetros.","Áreas de muro y vanos"),
("Dato faltante","Tipos de muro M1–M35 (etiquetas del plano) se definen en Xref_Muros_Sotano, no incluido. Se clasificó por capa y leyenda: rojo = carga, gris = no carga; doble hilada sobre MC1 = muro de contención.","Especificación de block y refuerzo"),
("Supuesto","Vanos detectados como huecos colineales de 0.55–2.50 m en muros de edificación (5 vanos). Puede haber vanos junto a esquinas no detectables sin plano arquitectónico.","Deducción de vanos"),
("Supuesto","Huecos 0.10–0.85 m entre tramos = columnas o mochetas (no muro): "+hs,"Se cuantifican en columnas/mochetas"),
("RFI","M-40, M-41, M-42 y M-49 (block que no carga, y = 7.81) corren paralelos al muro MC1 a 0.88 m, sobre su cimiento. Confirmar si forman parte del sistema MC1 (cámara) o son muros de edificación.","11.63 ml"),
("RFI","Remates de doble hilada al NE, fuera de Z1-d y Z1-e (MC-16, MC-21, MC-22, MC-23), se incluyeron como muro MC1.","3.14 ml"),
("Verificación","Longitud por eje vs. conteo de unidades (n × 0.40): MC1 −2.3 %, 0.15 carga −5.3 %, 0.20 carga −4.3 %, 0.15 no carga −10.4 %. Las unidades se dibujaron traslapadas (paso real 0.36–0.39 m); manda la longitud por eje.",""),
("Limitación","Los muros se entregan como bloques con atributos (doble clic) por tramo; no es el bloque dinámico de muro de proyectos anteriores (no está disponible aquí). Sustituir en AutoCAD si se requiere.","")]
for i,(t,d,im) in enumerate(items,1): row(wq,r,[i,t,d,im],left=(3,4)); wq.row_dimensions[r].height=42; r+=1
for w in wb.worksheets:
    w.page_setup.orientation="landscape"; w.sheet_properties.pageSetUpPr.fitToPage=True; w.page_setup.fitToWidth=1; w.page_setup.fitToHeight=0
wb.save(out); print("ok")
