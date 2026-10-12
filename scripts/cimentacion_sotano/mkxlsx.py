import sys, json, math
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
R=json.load(open(sys.argv[1])); out=sys.argv[2]
F=Font(name="Arial",size=10,color="000000"); FB=Font(name="Arial",size=10,bold=True,color="000000")
FT=Font(name="Arial",size=12,bold=True,color="000000"); FI=Font(name="Arial",size=9,italic=True,color="595959")
C=Alignment(horizontal="center",vertical="center",wrap_text=True); LFT=Alignment(horizontal="left",vertical="center",wrap_text=True)
HDR=PatternFill("solid",fgColor="D9D9D9"); GRP=PatternFill("solid",fgColor="F2F2F2"); INP=PatternFill("solid",fgColor="FFF2CC")
thin=Side(style="thin",color="BFBFBF"); B=Border(bottom=thin,top=thin,left=thin,right=thin)
wb=Workbook()
def sheet(ws,title,sub,headers,widths):
    ws.sheet_view.showGridLines=False
    ws["A1"]=title; ws["A1"].font=FT; ws["A2"]=sub; ws["A2"].font=FI
    for i,(h,w) in enumerate(zip(headers,widths),1):
        c=ws.cell(row=4,column=i,value=h); c.font=FB; c.fill=HDR; c.alignment=C; c.border=B
        ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes="A5"; return 5
def row(ws,r,vals,bold=False,fill=None,fmt=None,left=()):
    for i,v in enumerate(vals,1):
        c=ws.cell(row=r,column=i,value=v); c.font=FB if bold else F; c.alignment=LFT if i in left else C; c.border=B
        if fill: c.fill=fill
        if fmt and i in fmt: c.number_format=fmt[i]
SUB="Proyecto: Sótano – Plano S-1 Cimentación y detalles (DXF) | Fecha: 2026-10-12 | Unidades: m, m², m³"
# ---------- Parámetros ----------
wp=wb.active; wp.title="Parámetros"
r=sheet(wp,"PARÁMETROS DE SECCIÓN",SUB+" | Celdas amarillas = dato pendiente de detalle (xref no incluido)",["Parámetro","Valor","Unidad","Fuente / estado"],[42,12,10,58])
P={}
pars=[("Peralte zapata Z1","","m","PENDIENTE – Xref_Detalles_de_Cimetacion"),("Peralte zapata Z2","","m","PENDIENTE"),("Peralte zapata Z3","","m","PENDIENTE"),
("Peralte zapata Z4","","m","PENDIENTE"),("Peralte zapata Z5","","m","PENDIENTE"),("Peralte zapata Z6","","m","PENDIENTE"),("Peralte zapata Z7","","m","PENDIENTE"),("Peralte zapata Z8","","m","PENDIENTE"),
("Peralte cimiento corrido CC1","","m","PENDIENTE – detalle típico CC1"),("Peralte cimiento MC1","","m","PENDIENTE – detalle MC1"),("Altura muro MC1 (relleno)","","m","PENDIENTE – detalle MC1"),
("Peralte solera SA","","m","PENDIENTE – detalle SA"),("Espesor concreto pobre",0.04,"m","Dato – nota 1 detalle MC1"),("Espesor contrapiso CP",0.10,"m","Dato – leyenda CP")]
for n,v,u,f in pars:
    row(wp,r,[n,v,u,f],fill=None,fmt={2:"0.00"},left=(1,4))
    if v=="": wp.cell(row=r,column=2).fill=INP
    P[n]=f"Parámetros!$B${r}"; r+=1
# ---------- Zapatas ----------
wz=wb.create_sheet("Zapatas")
r=sheet(wz,"ZAPATAS – EXTRACCIÓN",SUB,["ID","Tipo","Largo","Ancho","Área","Peralte","Volumen","Centro X","Centro Y","Observaciones"],[9,8,9,9,10,10,11,11,11,46])
wz.cell(row=3,column=3,value="m").font=FI; wz.cell(row=3,column=5,value="m²").font=FI; wz.cell(row=3,column=7,value="m³").font=FI
z0=r
obs={"Z1-e":"Dibujada 2.50×2.50; las otras 4 Z1 miden 2.20×2.20 → RFI","Z5":"Cotas en plano 7.90 × 3.90","Z6":"Cuadrilátero no rectangular (lados 3.70–3.77)","Z7":"Cota en plano 5.75","Z8":"Cota en plano 6.35 vs 6.55 medido; contiene CC1 → RFI","Z3-a":"Lado norte bajo MC1 (línea oculta); se toma 3.30 como Z3 típica"}
for k,v in sorted(R["zap"].items(),key=lambda kv:(kv[1]["tipo"],kv[0])):
    l=v["lados"]; a=(l[0]+l[2])/2; b=(l[1]+l[3])/2; p=v["pts"]
    cx=sum(q[0] for q in p)/4; cy=sum(q[1] for q in p)/4
    row(wz,r,[k,v["tipo"],round(a,3),round(b,3),round(v["area"],3),f"={P['Peralte zapata '+v['tipo']]}" ,f'=IF(F{r}="","",E{r}*F{r})',round(cx,2),round(cy,2),obs.get(k,"")],fmt={3:"0.00",4:"0.00",5:"0.00",6:"0.00",7:"0.00",8:"0.00",9:"0.00"},left=(10,))
    wz.cell(row=r,column=6).value=f'=IF({P["Peralte zapata "+v["tipo"]]}="","",{P["Peralte zapata "+v["tipo"]]})'
    r+=1
z1=r-1
row(wz,r,["TOTAL","",None,None,f"=SUM(E{z0}:E{z1})","",f'=IF(COUNT(G{z0}:G{z1})<COUNTA(A{z0}:A{z1}),"pendiente",SUM(G{z0}:G{z1}))',None,None,f"{z1-z0+1} zapatas"],bold=True,fill=GRP,fmt={5:"0.00",7:"0.00"})
ZT=f"Zapatas!$E${r}"; ZV=f"Zapatas!$G${r}"; zrng=(z0,z1)
# ---------- Cimientos ----------
wc=wb.create_sheet("Cimientos")
r=sheet(wc,"CIMIENTOS CORRIDOS – EXTRACCIÓN",SUB,["Grupo","Elemento","Concepto","Unidad","Cantidad","Método","Observaciones"],[18,22,40,8,11,40,46])
mc=R["MC"]; cc=R["CC"]; pa=R["patches"]
rows=[
("Muro de contención","Cimiento MC1 1.00","Huella bruta (incluye tramos bajo zapatas)","m²",mc["area"]+mc["area_bajo_zap"],"Unión de franjas por pares de bordes + juntas","Verificación: semisuma de bordes = 111.86 ml (Δ 0.8 %)"),
("Muro de contención","Cimiento MC1 1.00","(−) Tramos bajo zapatas","m²",-mc["area_bajo_zap"],"Intersección huella MC1 ∩ zapatas","Se asigna a la zapata (cimiento termina en cara de zapata)"),
("Muro de contención","Cimiento MC1 1.00","Huella neta","m²",mc["area"],"Bruta − bajo zapatas","Cantidad de control para concreto y concreto pobre"),
("Muro de contención","Cimiento MC1 1.00","Longitud equivalente neta","ml",mc["area"]/1.0,"Huella neta / 1.00","Ejes trazados en DXF: %.2f ml (no incluyen juntas de esquina)"%mc["L_cl"]),
("Muro de contención","Cimiento MC1 1.00","Longitud equivalente bruta (muro continuo)","ml",(mc["area"]+mc["area_bajo_zap"])/1.0,"Huella bruta / 1.00","Base para muro, drenaje e impermeabilización"),
("Edificación","Cimiento corrido CC1 0.70","Huella bruta","m²",cc["area"]+cc["area_bajo_zap"],"Unión de franjas − MC1",""),
("Edificación","Cimiento corrido CC1 0.70","(−) Tramos bajo zapatas","m²",-cc["area_bajo_zap"],"Intersección CC1 ∩ zapatas","Incluye %.2f m² dentro de Z8 (= %.2f ml) → RFI"%(cc["area_en_Z8"],cc["area_en_Z8"]/0.7)),
("Edificación","Cimiento corrido CC1 0.70","Huella neta","m²",cc["area"],"Bruta − bajo zapatas","Cantidad de control para concreto"),
("Edificación","Cimiento corrido CC1 0.70","Longitud equivalente neta","ml",cc["area"]/0.7,"Huella neta / 0.70","Ejes en DXF: %.2f ml; diferencia = franjas de CC1 que sobresalen de Z2, Z3 y Z8"%cc["L_cl"]),
]
for g,e,c,u,q,mtd,o in rows:
    row(wc,r,[g,e,c,u,round(q,3),mtd,o],bold=c.startswith("Huella neta"),fill=GRP if c.startswith("Huella neta") else None,fmt={5:"0.00"},left=(3,6,7)); r+=1
r+=1; wc.cell(row=r,column=1,value="Juntas y esquinas cerradas con polígono (coordenadas en DXF)").font=FB; r+=1
for k,v in pa.items():
    row(wc,r,["Muro de contención" if k.startswith("MC1") else "Edificación",k.split(" ")[0],k,"m²",round(v["area"],3),"Esquina ext. + esquina int. + proyecciones","Incluido en la huella (unión, sin doble conteo)"],fmt={5:"0.000"},left=(3,6,7)); r+=1
CC_NET=None
# ---------- Soleras ----------
ws=wb.create_sheet("Soleras SA")
r=sheet(ws,"SOLERA SA 0.30 – EXTRACCIÓN",SUB,["Tramo","Longitud según trazo","Unidad","Observaciones"],[10,20,8,60])
s0=r
for i,t in enumerate(R["SA_tramos"],1):
    row(ws,r,[f"SA-{i:02d}",round(t["L"],3),"ml","De (%.2f, %.2f) a (%.2f, %.2f)"%(t["p"][0],t["p"][1],t["q"][0],t["q"][1])],fmt={2:"0.00"},left=(4,)); r+=1
row(ws,r,["TOTAL",f"=SUM(B{s0}:B{r-1})","ml","Trazo entre extremos dibujados (cruza zapatas hasta columna)"],bold=True,fill=GRP,fmt={2:"0.00"},left=(4,)); SAT=f"'Soleras SA'!$B${r}"; r+=1
row(ws,r,["Neto",round(R["SA"]["area"]/0.3,3),"ml","Fuera de huella de zapatas (huella neta %.3f m² / 0.30)"%R["SA"]["area"]],bold=True,fill=GRP,fmt={2:"0.00"},left=(4,)); SAN=f"'Soleras SA'!$B${r}"
# ---------- Resumen ----------
wr=wb.create_sheet("Resumen",0)
r=sheet(wr,"RESUMEN DE CUANTIFICACIÓN – CIMENTACIÓN SÓTANO",SUB,["No.","Elemento","Sección / tipo","Unidad","Cantidad","Método","Fuente","Observaciones"],[6,34,18,8,12,30,24,48])
def grp(t):
    global r
    for i in range(1,9): 
        c=wr.cell(row=r,column=i); c.fill=GRP; c.border=B
    wr.cell(row=r,column=1,value=t).font=FB; wr.merge_cells(start_row=r,start_column=1,end_row=r,end_column=8); wr.cell(row=r,column=1).alignment=LFT; r+=1
n=[0]
def it(e,s,u,q,m,f,o,fmt="0.00"):
    global r; n[0]+=1
    row(wr,r,[n[0],e,s,u,q,m,f,o],fmt={5:fmt},left=(2,6,8)); r+=1
grp("A. CIMENTACIÓN – EDIFICACIÓN")
from collections import OrderedDict
tipos=OrderedDict()
for k,v in sorted(R["zap"].items(),key=lambda kv:(kv[1]["tipo"],kv[0])):
    l=v["lados"]; key=(v["tipo"],round((l[0]+l[2])/2,2),round((l[1]+l[3])/2,2)); tipos.setdefault(key,[]).append(k)
for (t,a,b),ids in tipos.items():
    it(f"Zapata {t}",f"{a:.2f} × {b:.2f}","u",len(ids),"Conteo de bloques en DXF","Capa ZAPATA (planta S-1)",", ".join(ids)+(" – RFI medida" if t=="Z1" and a>2.3 else ""),fmt="0")
it("Zapatas – área en planta","Z1 a Z8","m²",f"={ZT}","Suma de áreas (hoja Zapatas)","Capa ZAPATA","")
it("Zapatas – concreto","Z1 a Z8","m³",f"={ZV}","Área × peralte","Detalle de zapatas (xref)","Pendiente: ingresar peraltes en hoja Parámetros")
it("Cimiento corrido CC1","0.70","ml",round(cc["area"]/0.7,2),"Huella neta / 0.70","Capa CIMIENTO 004","Excluye tramos bajo zapatas (incl. %.2f ml dentro de Z8 → RFI)"%(cc["area_en_Z8"]/0.7))
it("Cimiento corrido CC1 – huella","0.70","m²",round(cc["area"],2),"Unión de franjas − MC1 − zapatas","Capa CIMIENTO 004","")
q=round(cc["area"],3)
it("Cimiento corrido CC1 – concreto","0.70","m³",f'=IF({P["Peralte cimiento corrido CC1"]}="","pendiente",{q}*{P["Peralte cimiento corrido CC1"]})',"Huella × peralte","Detalle típico CC1 (xref)","Pendiente peralte")
it("Solera SA","0.30","ml",f"={SAT}","Suma de tramos según trazo","Capa S-solera de amarre 003","Fuera de zapatas: %.2f ml. Usar según nivel del detalle SA"%(R["SA"]["area"]/0.3))
cols={"C1":3,"C2":2,"C3":1,"C4":3,"C5":8,"C6":1,"C7":1}
it("Columnas (arranques sobre zapata)","C1–C7","u",sum(cols.values()),"Conteo de etiquetas","Textos TEX5",", ".join(f"{k}: {v}" for k,v in cols.items())+". Secciones en detalle de columnas (xref)",fmt="0")
it("Contrapiso CP","t = 0.10","m²",285.46,"Área neta de hatch (par-impar)","Capa S-contrapiso 001","")
it("Contrapiso CP – concreto","t = 0.10","m³",f"=ROUND(285.462*{P['Espesor contrapiso CP']},3)","Área × espesor","Leyenda CP","")
grp("B. MURO DE CONTENCIÓN MC1 (grupo independiente)")
it("Cimiento MC1","1.00","ml",round(mc["area"],2),"Huella neta / 1.00","Capa CIMIENTO 004","Excluye %.2f m² bajo zapatas (Z2, Z3, Z5, Z8, Z1)"%mc["area_bajo_zap"])
it("Cimiento MC1 – huella","1.00","m²",round(mc["area"],2),"Unión de franjas + 7 juntas","Capa CIMIENTO 004","Verificación independiente Δ 0.8 %")
q=round(mc["area"],3)
it("Cimiento MC1 – concreto","1.00","m³",f'=IF({P["Peralte cimiento MC1"]}="","pendiente",{q}*{P["Peralte cimiento MC1"]})',"Huella × peralte","Detalle MC1 (xref)","Pendiente peralte")
it("Concreto pobre bajo cimiento MC1","t = 0.04","m³",f"=ROUND({q}*{P['Espesor concreto pobre']},3)","Huella neta × 0.04","Nota 1, detalle MC1","[Supuesto] solo bajo cimiento MC1")
qb=round(mc["area"]+mc["area_bajo_zap"],2)
it("Muro MC1","—","ml",qb,"Huella bruta / 1.00 (continuo sobre zapatas)","Capa CIMIENTO 004","[Supuesto] el muro corre continuo; altura pendiente")
it("Muro MC1 – área","—","m²",f'=IF({P["Altura muro MC1 (relleno)"]}="","pendiente",{qb}*{P["Altura muro MC1 (relleno)"]})',"ml × altura","Detalle MC1 (xref)","Pendiente altura")
it("Drenaje francés tubo Ø 6\" perforado","—","ml",qb,"Igual a longitud de muro","Nota 3, detalle MC1","[Supuesto] a todo lo largo del muro")
it("Geocompuesto drenante (Mac Drain 1L)","—","m²",f'=IF({P["Altura muro MC1 (relleno)"]}="","pendiente",{qb}*{P["Altura muro MC1 (relleno)"]})',"ml × altura","Nota 4, detalle MC1","Pendiente altura")
it("Impermeabilización cara expuesta (Sikatop 144)","—","m²",f'=IF({P["Altura muro MC1 (relleno)"]}="","pendiente",{qb}*{P["Altura muro MC1 (relleno)"]})',"ml × altura","Nota 5, detalle MC1","Pendiente altura")
it("Pines (detalles DP)","DP2/DP3/DP4/DP5","—","ver obs.","Conteo de llamadas","Textos S-text 003","Llamadas: DP2 = 39, DP3 = 3, DP4 = 6, DP5 = 20. Cantidad de pines = separación de detalle (xref) → pendiente")
r+=1
wr.cell(row=r,column=1,value="Excavación, relleno y acero de refuerzo: pendientes del nivel de desplante y de los armados (xref de detalles no incluidos en el DXF). Ver hoja Pendientes.").font=FI
wr.merge_cells(start_row=r,start_column=1,end_row=r,end_column=8)
# ---------- Pendientes ----------
wq=wb.create_sheet("Pendientes")
r=sheet(wq,"PENDIENTES, RFI Y SUPUESTOS",SUB,["No.","Tipo","Descripción","Impacto"],[6,12,90,40])
pend=[
("Dato faltante","Los detalles vienen de xrefs NO incrustados en el DXF: Xref_Detalles_de_Cimetacion, Xref_detalle_de_pineado, Xref_Muros_Sotano, Xref_Ejes_Sotano. Sin ellos no hay peraltes, armados ni desplante.","Concreto, acero, excavación y relleno quedan con fórmula lista pero pendientes"),
("RFI","Z1 al NE (Z1-e) está dibujada de 2.50 × 2.50; las otras cuatro Z1 miden 2.20 × 2.20.","Δ 1.41 m² de área"),
("RFI","Z8 mide 6.55 × 7.09 m (cota en plano 6.35) y contiene cimientos CC1 (%.2f ml). Confirmar si Z8 es losa/zapata combinada y si los CC1 van sobre ella o se eliminan."%(cc["area_en_Z8"]/0.7),"CC1 ± %.2f ml; volumen Z8"%(cc["area_en_Z8"]/0.7)),
("RFI","Solera SA: confirmar si se mide entre caras de zapata (%.2f ml) o de columna a columna según trazo (%.2f ml)."%(R["SA"]["area"]/0.3,sum(t["L"] for t in R["SA_tramos"])),"Hasta 32 ml de diferencia"),
("Pendiente","Dos remates de cimiento de 0.80 m de ancho que salen de Z1-d y Z1-e (NE) no se cuantificaron (≈ 0.8 m² c/u).","Menor"),
("Supuesto","Intersecciones: zapata > MC1 > CC1. Cada cimiento termina en la cara de la zapata o del cimiento que lo cruza.","Sin dobles conteos"),
("Supuesto","Concreto pobre t = 0.04 (nota 1) solo bajo la huella neta del cimiento MC1.","Si va bajo zapatas: +%.2f m²"%mc["area_bajo_zap"]),
("Supuesto","Muro MC1, drenaje francés e impermeabilización a todo lo largo del MC1 (%.2f ml)."%qb,""),
("Limitación","Bordes de MC1 dibujados a 1.005–1.016 m en vez de 1.00 (imprecisión de dibujo); se usó el ancho dibujado.","< 1 %"),
("Limitación","Archivo entregado en DXF. Abrir en AutoCAD y guardar como DWG con el mismo nombre; crear allí los filtros de grupo (Cuantificación > Cimentación / Muro de contención) y ejecutar PURGE + AUDIT.",""),
]
for i,(t,dsc,imp) in enumerate(pend,1):
    row(wq,r,[i,t,dsc,imp],left=(3,4)); wq.row_dimensions[r].height=30; r+=1
for w in wb.worksheets:
    w.page_setup.orientation="landscape"; w.page_setup.fitToWidth=1; w.sheet_properties.pageSetUpPr.fitToPage=True; w.page_setup.fitToHeight=0
wb.save(out); print("ok",out)
