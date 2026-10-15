"""comparar.py – compara la cuantificación automática contra la hecha a mano (DXF y/o Excel).

python comparar.py --mio "<plano> - CUANT.dxf" --manual "<plano a mano>.dxf" [--excel "<extracción a mano>.xlsx"]
                   [--factor 1] [--salida comparacion.xlsx] [--config config.json]

Clases (se leen del nombre de capa con los patrones de config.json > comparacion):
  Muro <cm>cm, Vano <cm>cm, Cimiento <ancho> (ancho leído de "de 0.40 x 0.20").
Elementos leídos igual en los dos archivos:
  INSERT  -> largo = extensión de su geometría a lo largo de la rotación del bloque (sirve para xMuro dinámico y estático)
  polilínea cerrada -> rectángulo: el lado largo es el eje
  polilínea abierta -> cada tramo recto es un elemento (cimiento)
  LINE en capas de cuantificación -> no se suman (remates), se cuentan aparte
Faltante = elemento a mano cubierto < 50 % por elementos automáticos de su clase (paralelos, a <= dist_eje del eje).
Sobrante = lo mismo al revés. Diferencias por clase en m y %.
"""
import argparse
import json
import math
import os
import re
import sys

import ezdxf
from openpyxl import Workbook, load_workbook

from extraccion import FB, _fila, _hoja

PATRONES = {"vano": r"Muro de (\d+)\s*cm.*Vano", "muro": r"Muro de (\d+)\s*cm", "cimiento": r"Cimiento corrido"}


def clase_de_capa(capa, pat):
    if re.search(pat["vano"], capa, re.I):
        return f"Vano {int(re.search(pat['vano'], capa, re.I).group(1))}cm"
    m = re.search(pat["muro"], capa, re.I)
    if m:
        return f"Muro {int(m.group(1))}cm"
    if re.search(pat["cimiento"], capa, re.I):
        a = re.search(r"de\s+(\d+(?:[.,]\d+)?)\s*[x×]", capa, re.I)
        return f"Cimiento {float(a.group(1).replace(',', '.')):.2f}" if a else f"Cimiento ({capa})"
    return None


def _puntos(e):
    pts = []
    for v in e.virtual_entities():
        t = v.dxftype()
        if t == "LINE":
            pts += [v.dxf.start, v.dxf.end]
        elif t == "LWPOLYLINE":
            pts += [p for p in v.vertices_in_wcs()]
        elif t == "POLYLINE":
            pts += [p.dxf.location for p in v.vertices]
        elif t == "INSERT":
            pts += _puntos(v)
    return pts


def elementos(ruta, factor, pat):
    doc = ezdxf.readfile(ruta)
    out, lineas = [], 0
    for e in doc.modelspace():
        cl = clase_de_capa(e.dxf.layer, pat)
        if cl is None:
            continue
        t = e.dxftype()
        if t == "LINE":
            lineas += 1
            continue
        if t == "INSERT":
            pts = _puntos(e)
            if not pts:
                continue
            r = math.radians(e.dxf.rotation)
            u, v = (math.cos(r), math.sin(r)), (-math.sin(r), math.cos(r))
            pu = [p[0] * u[0] + p[1] * u[1] for p in pts]
            pv = [p[0] * v[0] + p[1] * v[1] for p in pts]
            a, b = min(pu), max(pu)
            cv = (min(pv) + max(pv)) / 2
            p0 = (u[0] * a + v[0] * cv, u[1] * a + v[1] * cv)
            p1 = (u[0] * b + v[0] * cv, u[1] * b + v[1] * cv)
            out.append(_el(cl, e, p0, p1, factor))
        elif t in ("LWPOLYLINE", "POLYLINE"):
            pts = [tuple(p)[:2] for p in (e.vertices_in_wcs() if t == "LWPOLYLINE" else [v.dxf.location for v in e.vertices])]
            cerrada = e.closed if t == "LWPOLYLINE" else e.is_closed
            if cerrada and len(pts) == 4:
                lados = [(pts[i], pts[(i + 1) % 4]) for i in range(4)]
                L0, L1 = math.dist(*lados[0]), math.dist(*lados[1])
                (a0, a1), (b0, b1) = (lados[0], lados[2]) if L0 >= L1 else (lados[1], lados[3])
                p0 = ((a0[0] + b1[0]) / 2, (a0[1] + b1[1]) / 2)
                p1 = ((a1[0] + b0[0]) / 2, (a1[1] + b0[1]) / 2)
                out.append(_el(cl, e, p0, p1, factor))
            else:
                if cerrada:
                    pts = pts + [pts[0]]
                for a, b in zip(pts, pts[1:]):
                    if math.dist(a, b) > 1e-9:
                        out.append(_el(cl, e, a, b, factor))
    return out, lineas


def _el(cl, e, p0, p1, f):
    p0 = (p0[0] * f, p0[1] * f)
    p1 = (p1[0] * f, p1[1] * f)
    return dict(clase=cl, capa=e.dxf.layer, handle=e.dxf.handle, p0=p0, p1=p1, largo=math.dist(p0, p1))


def cubierto(A, otros, dist_eje, ang):
    L = A["largo"]
    if L <= 0:
        return 0.0
    d = ((A["p1"][0] - A["p0"][0]) / L, (A["p1"][1] - A["p0"][1]) / L)
    n = (-d[1], d[0])
    ints = []
    smin = math.sin(math.radians(ang))
    for B in otros:
        if B["clase"] != A["clase"] or B["largo"] <= 0:
            continue
        db = ((B["p1"][0] - B["p0"][0]) / B["largo"], (B["p1"][1] - B["p0"][1]) / B["largo"])
        if abs(d[0] * db[1] - d[1] * db[0]) > smin:
            continue
        mb = ((B["p0"][0] + B["p1"][0]) / 2 - A["p0"][0], (B["p0"][1] + B["p1"][1]) / 2 - A["p0"][1])
        if abs(mb[0] * n[0] + mb[1] * n[1]) > dist_eje:
            continue
        t = sorted(((q[0] - A["p0"][0]) * d[0] + (q[1] - A["p0"][1]) * d[1] for q in (B["p0"], B["p1"])))
        a, b = max(0.0, t[0]), min(L, t[1])
        if b > a:
            ints.append((a, b))
    ints.sort()
    tot, fin = 0.0, -1.0
    for a, b in ints:
        if b <= fin:
            continue
        tot += b - max(a, fin)
        fin = b
    return min(tot / L, 1.0)


def excel_manual(ruta, pat, hoja=None):
    """Suma por clase de la extracción a mano. Busca encabezados Layer/Capa y Length/Longitud (o 'Longitud de Muro')."""
    wb = load_workbook(ruta, data_only=True, read_only=True)
    hojas = [wb[hoja]] if hoja else wb.worksheets
    tot, filas = {}, 0
    for ws in hojas:
        rows = list(ws.iter_rows(values_only=True))
        for i, r in enumerate(rows[:30]):
            h = [str(c).strip().lower() if c is not None else "" for c in r]
            ic = next((j for j, x in enumerate(h) if x in ("layer", "capa")), None)
            il = next((j for j, x in enumerate(h) if x in ("longitud de muro",)), None)
            il2 = next((j for j, x in enumerate(h) if x in ("length", "longitud")), None)
            icnt = next((j for j, x in enumerate(h) if x in ("count", "cantidad", "recuento")), None)
            if ic is None or (il is None and il2 is None):
                continue
            for r2 in rows[i + 1:]:
                if r2 is None or ic >= len(r2) or not r2[ic]:
                    continue
                cl = clase_de_capa(str(r2[ic]), pat)
                if cl is None:
                    continue
                v = None
                for j in (il, il2):
                    if j is not None and j < len(r2) and isinstance(r2[j], (int, float)):
                        v = float(r2[j])
                        break
                if v is None:
                    continue
                c = float(r2[icnt]) if icnt is not None and icnt < len(r2) and isinstance(r2[icnt], (int, float)) else 1.0
                tot[cl] = tot.get(cl, 0.0) + v * c
                filas += 1
            break
    return tot, filas


def comparar(mio, manual, salida, factor=1.0, excel=None, cfg=None):
    c = (cfg or {}).get("comparacion", {})
    pat = dict(PATRONES, **c.get("patrones", {}))
    dist_eje, ang = c.get("dist_eje", 0.10), c.get("tol_angular", 5.0)
    A, la = elementos(mio, factor, pat)
    M, lm = elementos(manual, factor, pat)
    for e in A:
        e["cob"] = cubierto(e, M, dist_eje, ang)
    for e in M:
        e["cob"] = cubierto(e, A, dist_eje, ang)
    ex, nfx = excel_manual(excel, pat) if excel else ({}, 0)
    clases = sorted({e["clase"] for e in A + M} | set(ex))

    wb = Workbook()
    wb.remove(wb.active)
    cols = [("Clase", 22), ("Longitud automática", 14), ("Longitud a mano (DXF)", 14), ("Diferencia DXF", 13), ("Diferencia DXF %", 12),
            ("Longitud a mano (Excel)", 14), ("Diferencia Excel", 13), ("Diferencia Excel %", 12), ("Unidad", 8),
            ("Objetos automáticos", 11), ("Objetos a mano", 11), ("Faltantes", 10), ("Sobrantes", 10)]
    ws = _hoja(wb, "RESUMEN", cols)
    res = {}
    for i, cl in enumerate(clases, start=2):
        la_ = sum(e["largo"] for e in A if e["clase"] == cl)
        lm_ = sum(e["largo"] for e in M if e["clase"] == cl)
        falt = sum(1 for e in M if e["clase"] == cl and e["cob"] < 0.5)
        sobr = sum(1 for e in A if e["clase"] == cl and e["cob"] < 0.5)
        lx = ex.get(cl)
        _fila(ws, i, [cl, la_, lm_, f"=B{i}-C{i}", f'=IF(C{i}=0,"",(B{i}-C{i})/C{i})',
                      lx, f'=IF(F{i}="","",B{i}-F{i})' if lx is not None else None,
                      f'=IF(OR(F{i}="",F{i}=0),"",(B{i}-F{i})/F{i})' if lx is not None else None, "m",
                      sum(1 for e in A if e["clase"] == cl), sum(1 for e in M if e["clase"] == cl), falt, sobr],
              {2: "0.000", 3: "0.000", 4: "0.000", 5: "0.00%", 6: "0.000", 7: "0.000", 8: "0.00%"})
        res[cl] = dict(auto=la_, manual=lm_, excel=lx, faltantes=falt, sobrantes=sobr)
    for nombre, lista, otro in (("FALTANTES", M, "automático"), ("SOBRANTES", A, "a mano")):
        w2 = _hoja(wb, nombre, [("Clase", 22), ("Capa", 40), ("Handle", 10), ("Longitud (m)", 12), ("Cubierto por " + otro, 14),
                                ("X inicio (m)", 12), ("Y inicio (m)", 12), ("X fin (m)", 12), ("Y fin (m)", 12)])
        r = 2
        for e in lista:
            if e["cob"] < 0.5:
                _fila(w2, r, [e["clase"], e["capa"], e["handle"], e["largo"], e["cob"], *e["p0"], *e["p1"]],
                      {4: "0.000", 5: "0.0%", 6: "0.000", 7: "0.000", 8: "0.000", 9: "0.000"})
                r += 1
    wd = _hoja(wb, "DETALLE", [("Origen", 12), ("Clase", 22), ("Capa", 40), ("Handle", 10), ("Longitud (m)", 12), ("Cubierto por el otro", 14),
                               ("X inicio (m)", 12), ("Y inicio (m)", 12), ("X fin (m)", 12), ("Y fin (m)", 12)])
    r = 2
    for org, lista in (("automático", A), ("a mano", M)):
        for e in sorted(lista, key=lambda e: (e["clase"], e["p0"])):
            _fila(wd, r, [org, e["clase"], e["capa"], e["handle"], e["largo"], e["cob"], *e["p0"], *e["p1"]],
                  {5: "0.000", 6: "0.0%", 7: "0.000", 8: "0.000", 9: "0.000", 10: "0.000"})
            r += 1
    wp = _hoja(wb, "PARAMETROS", [("Parámetro", 30), ("Valor", 80)])
    for i, (k, v) in enumerate([("DXF automático", mio), ("DXF a mano", manual), ("Excel a mano", excel or "(no)"),
                                ("Filas leídas del Excel", nfx), ("Factor a metros", factor), ("Distancia máxima entre ejes (m)", dist_eje),
                                ("Tolerancia angular (°)", ang), ("LINE en capas (no sumadas) automático / a mano", f"{la} / {lm}"),
                                ("Patrones de capa", json.dumps(pat, ensure_ascii=False))], start=2):
        _fila(wp, i, [k, v])
    wb.save(salida)
    return res


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--mio", required=True)
    ap.add_argument("--manual", required=True)
    ap.add_argument("--excel")
    ap.add_argument("--factor", type=float, default=1.0)
    ap.add_argument("--salida", default="comparacion.xlsx")
    ap.add_argument("--config")
    a = ap.parse_args(argv)
    cfg = json.load(open(a.config, encoding="utf-8")) if a.config else {}
    res = comparar(a.mio, a.manual, a.salida, a.factor, a.excel, cfg)
    for cl, r in res.items():
        d = r["auto"] - r["manual"]
        pct = f"{100 * d / r['manual']:.2f} %" if r["manual"] else "—"
        print(f"{cl:24} auto {r['auto']:10.3f}  mano {r['manual']:10.3f}  dif {d:+.3f} m ({pct})  faltan {r['faltantes']}  sobran {r['sobrantes']}")
    print(f"OK: {a.salida}")


if __name__ == "__main__":
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    main()
