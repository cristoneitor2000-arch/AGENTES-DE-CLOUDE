"""comparar.py contra una cuantificación a mano sintética con diferencias puestas a propósito."""
import copy

import openpyxl
import pytest

import comparar
import crear_manual_prueba
import pipeline


def test_comparacion(plano, conf, tipos_prueba, tmp_path):
    conv, cfg = conf
    info = pipeline.cuantificar(plano, ["A-MURO"], [.10, .15, .20], "N1", None, None, conv, cfg, log=lambda m: None,
                                carpeta_salida=str(tmp_path), tipos_cimiento=tipos_prueba)
    man, xls = str(tmp_path / "manual.dxf"), str(tmp_path / "manual.xlsx")
    crear_manual_prueba.crear(man, xls)
    out = str(tmp_path / "comparacion.xlsx")
    r = comparar.comparar(info["dxf"], man, out, 1.0, xls, cfg)
    # a mano la L va a ejes (7.00); con la regla del motor sin mocheta la esquina queda fuera (6.85)
    assert r["Muro 15cm"]["auto"] == pytest.approx(41.725) and r["Muro 15cm"]["manual"] == pytest.approx(43.875)
    assert r["Muro 15cm"]["faltantes"] == 1 and r["Muro 15cm"]["sobrantes"] == 0
    assert r["Vano 15cm"]["sobrantes"] == 1 and r["Vano 15cm"]["manual"] == pytest.approx(0.9)
    assert r["Cimiento 0.40"]["faltantes"] == 1
    for k in ("Muro 20cm", "Cimiento 0.50"):
        assert r[k]["auto"] == pytest.approx(r[k]["manual"]) and r[k]["faltantes"] == r[k]["sobrantes"] == 0
    assert r["Muro 10cm"]["auto"] == pytest.approx(2.95) and r["Muro 10cm"]["manual"] == pytest.approx(2.90)
    assert r["Muro 15cm"]["excel"] == pytest.approx(43.875)
    wb = openpyxl.load_workbook(out)
    f = list(wb["FALTANTES"].iter_rows(min_row=2, values_only=True))
    assert {(x[0], round(x[3], 3)) for x in f} == {("Muro 15cm", 2.0), ("Cimiento 0.40", 2.0)}
    s = list(wb["SOBRANTES"].iter_rows(min_row=2, values_only=True))
    assert [(x[0], round(x[3], 3)) for x in s] == [("Vano 15cm", 0.925)]
